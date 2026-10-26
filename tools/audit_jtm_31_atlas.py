#!/usr/bin/env python3
"""Read-only cross-reference audit for the 31-cancer JTM Word documents."""

from __future__ import annotations

import os
from pathlib import Path
import re

from docx import Document


folder = Path(os.environ.get("JTM_DIR", Path.home() / "Desktop/JTM"))
main = Document(folder / "manuscript_JTM_31_cancer_atlas.docx")
supp = Document(folder / "supplementary_material_JTM_31_cancer_atlas.docx")


def caption_numbers(document, kind: str, supplementary: bool) -> list[int]:
    token = "S" if supplementary else ""
    pattern = re.compile(rf"^{kind} {token}(\d+)\.")
    return [int(match.group(1)) for paragraph in document.paragraphs
            if (match := pattern.match(paragraph.text))]


def first_citations(document, kind: str, supplementary: bool) -> list[int]:
    prefix = "Supplementary " if supplementary else ""
    token = "S" if supplementary else ""
    pattern = re.compile(rf"{prefix}{kind} {token}(\d+)\b")
    seen: list[int] = []
    for paragraph in document.paragraphs:
        if re.match(r"^(?:Supplementary )?(?:Table|Figure) S?\d+\.", paragraph.text):
            continue
        for match in pattern.finditer(paragraph.text):
            number = int(match.group(1))
            if number not in seen:
                seen.append(number)
    return seen


for label, captions, citations in (
    ("main tables", caption_numbers(main, "Table", False),
     first_citations(main, "Table", False)),
    ("main figures", caption_numbers(main, "Figure", False),
     first_citations(main, "Figure", False)),
    ("supplementary tables", caption_numbers(supp, "Table", True),
     first_citations(main, "Table", True)),
    ("supplementary figures", caption_numbers(supp, "Figure", True),
     first_citations(main, "Figure", True)),
):
    expected = list(range(1, len(captions) + 1))
    if captions != expected or citations != expected:
        raise SystemExit(
            f"{label}: captions={captions}, first text citations={citations}; "
            f"expected {expected}"
        )
    print(f"{label}: {len(expected)} captions and citations in order")

all_text = "\n".join(p.text for d in (main, supp) for p in d.paragraphs)
if re.search(r"\bTable S?\d+[a-zA-Z]\b", all_text):
    raise SystemExit("A lettered table remains")
if re.search(r"Additional[ _-]file[ _-]4|machine-readable results", all_text, re.I):
    raise SystemExit("Additional file 4 or machine-readable-results wording remains")
if "cancer-endpoint" in all_text:
    raise SystemExit("Inconsistent cancer-endpoint terminology remains")

body = "\n".join(p.text for p in main.paragraphs[:132])
first_cited: list[int] = []
for match in re.finditer(r"\[([\d,\s\-–]+)\]", body):
    for item in match.group(1).split(","):
        item = item.strip()
        if "-" in item or "–" in item:
            left, right = re.split(r"[-–]", item, maxsplit=1)
            numbers = range(int(left), int(right) + 1)
        else:
            numbers = (int(item),)
        for number in numbers:
            if number not in first_cited:
                first_cited.append(number)
bibliography = [int(m.group(1)) for p in main.paragraphs
                if (m := re.match(r"^(\d+)\.\s", p.text))]
if bibliography != list(range(1, len(bibliography) + 1)):
    raise SystemExit("Bibliography numbering is incomplete")
if first_cited != bibliography:
    raise SystemExit(f"Literature first-citation order {first_cited} differs from bibliography {bibliography}")
print(f"literature: {len(bibliography)} references cited in order")
print("Additional file 4 and lettered tables: absent")
