#!/usr/bin/env python3
"""Audit the synchronized Journal of Translational Medicine submission files."""
from __future__ import annotations

import re
import csv
from zipfile import ZipFile
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "manuscript"
TABLES = ROOT / "results" / "tables"
CURRENT = {
    "manuscript_JTM_multifoundation_atlas.docx",
    "supplementary_material_JTM.docx",
    "response_to_reviewer_JTM.docx",
    "cover_letter_JTM.docx",
}


def require(condition: bool, message: str) -> None:
    if not condition:
        raise RuntimeError(message)


def text(document: Document) -> str:
    chunks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def csv_rows(name: str) -> list[dict[str, str]]:
    with (TABLES / name).open(newline="", encoding="utf-8") as source:
        return list(csv.DictReader(source))


for name in CURRENT:
    require((OUT / name).exists(), f"Missing current document: {name}")

main = Document(OUT / "manuscript_JTM_multifoundation_atlas.docx")
supp = Document(OUT / "supplementary_material_JTM.docx")
response = Document(OUT / "response_to_reviewer_JTM.docx")
cover = Document(OUT / "cover_letter_JTM.docx")
main_text = text(main)
supp_text = text(supp)
response_text = text(response)
cover_text = text(cover)
combined = "\n".join((main_text, supp_text, response_text, cover_text))

abstract_start = next(i for i, p in enumerate(main.paragraphs)
                      if p.text.strip() == "Abstract")
abstract_end = next(i for i, p in enumerate(main.paragraphs)
                    if p.text.startswith("Keywords:"))
abstract_text = " ".join(p.text for p in main.paragraphs[
    abstract_start + 1:abstract_end
])
abstract_words = re.findall(r"\b[\w²-]+\b", abstract_text)
require(len(abstract_words) <= 350,
        f"The structured abstract exceeds the journal's 350-word limit: {len(abstract_words)}")
for heading in ("Background.", "Methods.", "Results.", "Conclusions."):
    require(heading in abstract_text, f"Abstract section is missing: {heading}")
with ZipFile(OUT / "manuscript_JTM_multifoundation_atlas.docx") as archive:
    manuscript_xml = archive.read("word/document.xml")
require(b'w:type="page"' not in manuscript_xml and b"w:pageBreakBefore" not in manuscript_xml,
        "The main manuscript contains a manual page break")

require("Patient-level comparison of three released pathology representation pipelines" in main_text,
        "Current title is missing")
require("8,241" in main_text and "3,389" in main_text and "30 cancers" in main_text,
        "Current matched-cohort counts are missing")
require("1 to 20" in main_text or "1-to-20" in main_text,
        "Current 1 to 20 component range is missing")
require("TITAN" in main_text and "Giga-SSL" in main_text and "Prov-GigaPath" in main_text,
        "All three representations are not named")
require("PathoFMPred" in main_text and "PathoFMPred" in supp_text,
        "PathoFMPred naming is inconsistent")
require("DINOv2" in main_text, "Foundation-model introduction does not mention DINOv2")
require("TCGA-AA-A01F" in main_text and "TCGA-A6-A56B" in main_text,
        "The two COAD examples are missing from the main manuscript")
require("reference rank, not probability" in combined.lower(),
        "Binary rank warning is missing")
require("TITANPred" not in combined, "Obsolete TITANPred name remains")
require("revision-added" not in main_text.lower(), "Revision-history language remains in the manuscript")
require("inferentially qualified" not in main_text.lower(), "Overstated inferential terminology remains")
require("1,933" not in main_text and "1,507" not in main_text,
        "Stale matched-atlas task counts remain in the manuscript")
require("1,933" not in supp_text and "2,073" not in supp_text,
        "Stale task-universe counts remain in the Supplement")
require("1,933" not in response_text and "2,073" not in response_text,
        "Stale task-universe counts remain in the response")
require("1-10-component" not in combined and "1–10-component" not in combined,
        "Stale 1 to 10 primary component wording remains")
comparison = csv_rows("foundation_model_target_comparison.csv")
dictionary = csv_rows("endpoint_dictionary.csv")
provenance = {
    (row["outcome_type"], row["family"], row["tumor_type"],
     row["endpoint"], row["source"]): row
    for row in dictionary
}
direct = [
    row for row in comparison
    if row["outcome_type"] == "binary" and provenance[
        (row["outcome_type"], row["family"], row["tumor_type"],
         row["endpoint"], row["source"])
    ]["measurement_class"] == "directly observed genomic alteration"
]
direct_union = sum(int(row["supported_by_n"]) >= 1 for row in direct)
direct_all = sum(int(row["supported_by_n"]) == 3 for row in direct)
require(f"{direct_union}/{len(direct)}" in main_text and
        f"{direct_all}/{len(direct)}" in main_text,
        "Direct-genomic breadth counts do not match the refreshed atlas")
crossmodal_union = sum(
    int(row["supported_by_n"]) >= 1 and provenance[
        (row["outcome_type"], row["family"], row["tumor_type"],
         row["endpoint"], row["source"])
    ]["same_histology_modality"] == "FALSE"
    for row in comparison
)
require("Figure 3. Foundation-model leadership" in main_text and
        f"{crossmodal_union:,} cross-modal" in main_text,
        "Foundation-model leadership scope is not synchronized")
same_histology = {row["foundation_model"]: row for row in
                  csv_rows("foundation_model_same_histology_sensitivity.csv")}
expected_continuous = (
    same_histology["TITAN"]["cross_modal_continuous_crossings"] + ", " +
    same_histology["GigaSSL"]["cross_modal_continuous_crossings"] + " and " +
    same_histology["ProvGigaPath"]["cross_modal_continuous_crossings"]
)
require(expected_continuous in main_text,
        "Abstract representation order or continuous counts are inconsistent")
coad_binary = csv_rows("coad_pathofmpred_multifoundation_binary_predictions.csv")
coad_counts = {
    model: len({(row["family"], row["endpoint"]) for row in coad_binary
                if row["patient_id"] == "TCGA-AA-A01F" and
                row["foundation_model"] == model})
    for model in ("TITAN", "GigaSSL", "ProvGigaPath")
}
for model, count in coad_counts.items():
    display = {"GigaSSL": "Giga-SSL", "ProvGigaPath": "Prov-GigaPath"}.get(model, model)
    require(f"{display} returned {count}" in main_text,
            f"COAD binary model count is not synchronized for {model}")
require("GPL-3.0-or-later companion analysis repository" in supp_text,
        "Analysis-repository license statement is missing")
require("PathoFMPred contributor-authored source code and documentation are released under the MIT License" in supp_text,
        "PathoFMPred MIT source-code statement is missing")
require("TITAN fitted collection is excluded from the public repository" in supp_text,
        "TITAN redistribution boundary is missing")
require("—" not in combined, "Em dash remains in the submission documents")


def section_paragraphs(document: Document, start: str, end: str) -> list[str]:
    values = []
    active = False
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == start:
            active = True
            continue
        if paragraph.text.strip() == end:
            break
        if active:
            values.append(paragraph.text)
    return values


main_methods = "\n".join(section_paragraphs(main, "Methods", "Results"))
supp_methods = "\n".join(section_paragraphs(supp, "Supplementary Methods", "Table S1. Analysis coverage"))
require(not re.search(r"\b(?:We|we|Our|our)\b", main_methods),
        "Active first-person wording remains in the main Methods")
require(not re.search(r"\b(?:We|we|Our|our)\b", supp_methods),
        "Active first-person wording remains in the Supplementary Methods")


def parse_citations(value: str, maximum: int) -> list[int]:
    found = []
    for raw in re.findall(r"\[([0-9][0-9,\- ]*)\]", value):
        group = []
        for token in raw.split(","):
            token = token.strip()
            if "-" in token:
                start, finish = map(int, token.split("-", 1))
                group.extend(range(start, finish + 1))
            else:
                group.append(int(token))
        if group and all(number <= maximum for number in group):
            found.extend(group)
    return found


reference_heading = next(i for i, p in enumerate(main.paragraphs) if p.text.strip() == "References")
reference_numbers = []
for paragraph in main.paragraphs[reference_heading + 1:]:
    match = re.match(r"^(\d+)\.\s", paragraph.text)
    if match:
        reference_numbers.append(int(match.group(1)))
require(reference_numbers == list(range(1, len(reference_numbers) + 1)),
        "Bibliography is not consecutively numbered")
first_appearance = []
for paragraph in main.paragraphs[:reference_heading]:
    for number in parse_citations(paragraph.text, len(reference_numbers)):
        if number not in first_appearance:
            first_appearance.append(number)
for paragraph in supp.paragraphs:
    for number in parse_citations(paragraph.text, len(reference_numbers)):
        if number not in first_appearance:
            first_appearance.append(number)
for table in supp.tables:
    for row in table.rows:
        for cell in row.cells:
            for number in parse_citations(cell.text, len(reference_numbers)):
                if number not in first_appearance:
                    first_appearance.append(number)
require(first_appearance == reference_numbers,
        f"References are not numbered by first appearance: {first_appearance}")

table_captions = [p.text for p in main.paragraphs if re.match(r"^Table \d+\.", p.text)]
figure_captions = [p.text for p in main.paragraphs if re.match(r"^Figure \d+\.", p.text)]
require([int(re.match(r"Table (\d+)", x).group(1)) for x in table_captions] == list(range(1, len(table_captions) + 1)),
        "Main table captions are not sequential")
require([int(re.match(r"Figure (\d+)", x).group(1)) for x in figure_captions] == list(range(1, len(figure_captions) + 1)),
        "Main figure captions are not sequential")

first_supp = []
for paragraph in main.paragraphs:
    for label in re.findall(r"Supplementary Table S(\d+)", paragraph.text):
        if label not in first_supp:
            first_supp.append(label)
require(first_supp == sorted(first_supp, key=int),
        f"Supplementary tables are first cited out of order: {first_supp}")

for report in (
    "Additional_file_2_COAD_example_A_PathoFMPred_report.pdf",
    "Additional_file_3_COAD_example_B_PathoFMPred_report.pdf",
):
    require((OUT / report).exists() and (OUT / report).stat().st_size > 10_000,
            f"Missing or empty additional report: {report}")

print("Document audit passed for the synchronized main manuscript, Supplement, response, cover letter and patient reports.")
