#!/usr/bin/env python3
"""Audit the current matched AUROC/Q2 benchmark and its manuscript claims.

Denominators are derived from regenerated results, not an older fastPLS build.
"""

from __future__ import annotations

import csv
import math
import re
from collections import defaultdict
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
TABLES = ROOT / "results" / "tables"
DOCS = ROOT / "manuscript"
MODELS = ("TITAN", "GigaSSL", "ProvGigaPath")


def require(condition: bool, message: str) -> None:
    if not condition:
        raise RuntimeError(message)


def rows(name: str) -> list[dict[str, str]]:
    with (TABLES / name).open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def document_text(name: str) -> tuple[Document, str]:
    document = Document(DOCS / name)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return document, "\n".join(parts)


def truth(value: str) -> bool:
    require(value.upper() in {"TRUE", "FALSE"}, f"Not a Boolean: {value!r}")
    return value.upper() == "TRUE"


matched = rows("foundation_model_matched_screen.csv")
comparison = rows("foundation_model_target_comparison.csv")
require(matched and comparison, "Matched results or comparison table is empty")

key_columns = ("outcome_type", "family", "tumor_type", "endpoint", "source")
per_task: dict[tuple[str, ...], dict[str, dict[str, str]]] = defaultdict(dict)
for row in matched:
    model = row["foundation_model"]
    require(model in MODELS, f"Unknown representation: {model}")
    key = tuple(row[column] for column in key_columns)
    require(model not in per_task[key], f"Duplicate representation-task row: {model} {key}")
    per_task[key][model] = row
    require(row["fastPLS_version"] == "0.3", f"Incorrect fastPLS version: {model} {key}")
    require(row["fastPLS_remote_sha"] in {"", "NA"},
            f"A Git-built fastPLS row remains: {model} {key}")
    require(1 <= float(row["selected_components_min"]) <= 20,
            f"Invalid selected-component minimum: {model} {key}")
    require(1 <= float(row["selected_components_max"]) <= 20,
            f"Invalid selected-component maximum: {model} {key}")
    require(row["numerical_failure_or_fallback"].upper() == "FALSE",
            f"Numerical fallback or failure was recorded: {model} {key}")
    if row["outcome_type"] == "binary":
        auc = float(row["auc"])
        pr_auc = float(row["pr_auc"])
        prevalence = float(row["prevalence"])
        require(all(math.isfinite(value) and 0 <= value <= 1
                    for value in (auc, pr_auc, prevalence)),
                f"Invalid binary discrimination metric: {model} {key}")
        require(abs(float(row["no_skill_pr_auc"]) - prevalence) < 1e-10,
                f"The PR-AUC reference differs from prevalence: {model} {key}")
        require(int(row["positive"]) + int(row["negative"]) == int(row["n"]),
                f"Binary class counts do not sum to n: {model} {key}")
        require(row["primary_binary_metric"].upper().endswith("AUROC") and
                "AUROC" in row["primary_binary_tuning_objective"].upper(),
                f"Binary metric and tuning objective are misaligned: {model} {key}")
    else:
        require(math.isfinite(float(row["q2"])),
                f"Non-finite continuous Q2: {model} {key}")

require(all(set(models) == set(MODELS) for models in per_task.values()),
        "The matched task sets differ across the three representations")
require(len(comparison) == len(per_task),
        "Comparison rows do not match the unique matched tasks")

seen_comparison = set()
for row in comparison:
    key = tuple(row[column] for column in key_columns)
    require(key in per_task and key not in seen_comparison,
            f"Missing or duplicate comparison task: {key}")
    seen_comparison.add(key)
    count = 0
    for model in MODELS:
        source = per_task[key][model]
        effect = float(source["auc"] if row["outcome_type"] == "binary" else source["q2"])
        threshold = 0.60 if row["outcome_type"] == "binary" else 0.20
        require(abs(float(row[f"effect_{model}"]) - effect) < 1e-10,
                f"Comparison effect differs from matched screen: {model} {key}")
        crossing = truth(row[f"screening_positive_{model}"])
        require(crossing == (effect >= threshold),
                f"Crossing does not use the declared effect threshold: {model} {key}")
        count += crossing
    require(count == int(row["supported_by_n"]),
            f"Three-representation support count is inconsistent: {key}")

main, main_text = document_text("manuscript_JTM_multifoundation_atlas.docx")
supplement, supplement_text = document_text("supplementary_material_JTM.docx")
combined = main_text + "\n" + supplement_text
for phrase in ("TITAN", "Giga-SSL", "Prov-GigaPath", "PathoFMPred", "fastPLS 0.3"):
    require(phrase in main_text, f"Main manuscript omits {phrase}")
require(f"{len(per_task):,} cancer-endpoint" in main_text,
        "The matched task denominator is not synchronized in the manuscript")
require("1 to 20" in main_text or "1-to-20" in main_text,
        "The manuscript does not describe the 1-to-20-component probe")
require("AUROC" in main_text and "PR-AUC" in main_text,
        "Binary primary and precision-recall metrics are missing")
require("not probability" in combined.lower(),
        "Binary research-software ranks lack the non-probability warning")
for forbidden in ("TITANPred", "revision-added", "—"):
    require(forbidden not in combined, f"Obsolete or unwanted wording remains: {forbidden}")

for prefix in ("Table", "Figure"):
    numbers = [int(match.group(1)) for paragraph in main.paragraphs
               if (match := re.match(rf"^{prefix} (\d+)\. ", paragraph.text))]
    require(numbers and numbers == list(range(1, len(numbers) + 1)),
            f"Main {prefix.lower()} captions are not sequential: {numbers}")
require(len(main.inline_shapes) >= 3, "Main figures are missing from the Word document")
require(len(supplement.inline_shapes) >= 3,
        "Supplementary figures are missing from the Word document")
for name in (
    "Additional_file_2_COAD_example_A_PathoFMPred_report.pdf",
    "Additional_file_3_COAD_example_B_PathoFMPred_report.pdf",
):
    require((DOCS / name).is_file() and (DOCS / name).stat().st_size > 10_000,
            f"Additional report is missing or empty: {name}")
    require(name in combined, f"Additional report is not cited: {name}")

print(f"AUROC-centred release audit passed: {len(per_task):,} matched tasks, "
      f"{len(matched):,} representation-task rows, synchronized Word figures and reports")
