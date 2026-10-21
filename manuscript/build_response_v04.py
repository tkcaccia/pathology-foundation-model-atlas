#!/usr/bin/env python3
"""Write a concise response aligned with the current matched atlas."""

from pathlib import Path
import csv

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT = Path(__file__).resolve().parents[1]
TABLES = ROOT / "results" / "tables"
OUTPUT = ROOT / "manuscript" / "response_to_reviewer_JTM.docx"


def read(name):
    with (TABLES / name).open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def group_counts(rows, model, outcome):
    selected = [row for row in rows if row["foundation_model"] == model
                and row["outcome_type"] == outcome]
    metric = "q2" if outcome == "continuous" else "auc"
    threshold = 0.20 if outcome == "continuous" else 0.60
    crossing = sum(float(row[metric]) >= threshold for row in selected)
    return crossing, len(selected)


def add_section(document, title, prose):
    document.add_heading(title, level=1)
    for paragraph in prose:
        document.add_paragraph(paragraph)


def main():
    screen = read("foundation_model_matched_screen.csv")
    if not screen or any(row["fastPLS_version"] != "0.4" for row in screen):
        raise RuntimeError("The current fastPLS 0.4 benchmark is required")
    groups = {model: {kind: group_counts(screen, model, kind)
                      for kind in ("continuous", "binary")}
              for model in ("TITAN", "GigaSSL", "ProvGigaPath")}
    task_count = sum(1 for row in screen if row["foundation_model"] == "TITAN")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name in ("Title", "Heading 1"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = None
    doc.add_heading("Response to methodological review", 0)
    doc.add_paragraph(
        "Patient-level comparison of three released pathology slide "
        "representations with a common PLS-based probe across TCGA cancers"
    )
    doc.add_paragraph(
        "We revised the analysis around one matched comparison. The same "
        "patients, slide identifiers, outcome labels and patient-level folds "
        "were used for TITAN, Giga-SSL and Prov-GigaPath. We distinguish the "
        "descriptive benchmark from statistical discovery and clinical validation."
    )

    add_section(doc, "Matched inputs and analysis scope", [
        f"The primary comparison contains {task_count:,} cancer-endpoint pairs "
        "from 8,241 patients and 10,165 slide identifiers shared by all three "
        "representations. The full atlas retains H&E-derived TIL Regional "
        "Fraction as a same-modality outcome. We omit its 11 eligible pairs "
        "only from the default summary of reference phenotypes derived "
        "independently of H&E. A separate sensitivity uses all available "
        "slides for the same patients.",
        "Outcomes from published genomic, RNA, methylation and composite "
        "resources are linked at the participant or available sample-barcode "
        "level. Exact barcode agreement does not prove that the slide and assay "
        "sampled the same block or tumour region."
    ])
    add_section(doc, "Binary estimand and software version", [
        "The complete matched benchmark was refitted with fastPLS 0.4 and a "
        "1 to 20 component grid. Continuous components were selected by inner "
        "out-of-fold prediction error, with outer Q² as the primary result. "
        "Binary components were selected by inner out-of-fold AUROC, with "
        "outer AUROC as the primary result. Binary class calls use a threshold "
        "chosen only from the training data; PR-AUC and operating-point "
        "statistics are secondary.",
        "We fixed R's random-number algorithm before fold construction. This "
        "removes a sequential-versus-parallel seed discrepancy discovered "
        "during the rerun; all affected results were regenerated from the "
        "same documented analysis settings."
    ])
    totals = []
    for model, label in (("TITAN", "TITAN"), ("GigaSSL", "Giga-SSL"),
                         ("ProvGigaPath", "Prov-GigaPath")):
        continuous = groups[model]["continuous"]
        binary = groups[model]["binary"]
        totals.append(f"{label}: {continuous[0]}/{continuous[1]} continuous "
                      f"and {binary[0]}/{binary[1]} binary")
    add_section(doc, "Interpretation of the benchmark", [
        "The matched effect-threshold crossings are descriptive. At Q² at "
        "least 0.20 or AUROC at least 0.60, the counts are " + "; ".join(totals) + ". "
        "They are not three representation-specific FDR discovery counts, "
        "and correlated endpoint definitions mean they are not independent "
        "biological discoveries.",
        "The manuscript leads with feature-provenance classes, paired effects "
        "and named cancer-specific results. Alternative folds, tissue-source-site-"
        "code grouping and a ridge sensitivity show where apparent differences "
        "depend on the partition or downstream probe. Grouping is an internal "
        "cohort-structure test, not institutional or scanner validation."
    ])
    add_section(doc, "Research software and access", [
        "PathoFMPred is an R package for constructing and applying compatible "
        "linear research models. The public package contains no fitted TCGA "
        "collections. Separate access-controlled objects hold models for the "
        "three representations, subject to their respective upstream terms. "
        "The two COAD profiles illustrate the interface; the patients may "
        "have contributed labels to full-data model fitting, so the reports "
        "do not provide held-out case validation."
    ])
    add_section(doc, "Translational limits", [
        "No independent cohort was available. The manuscript therefore "
        "presents internally derived TCGA associations and research priorities, "
        "not assay replacement, calibrated clinical probabilities or evidence "
        "of transportability. The Supplement and machine-readable companions "
        "provide source definitions, missingness, fold adequacy and the full "
        "target-level results."
    ])

    for paragraph in doc.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
